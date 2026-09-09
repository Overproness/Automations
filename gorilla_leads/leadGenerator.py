import time
# import undetected_chromedriver as uc
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys  # Import Keys to send special keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
import openpyxl
import pyautogui

# Load Excel data
excel_file = "leads.xlsx"
workbook = openpyxl.load_workbook(excel_file)
sheet = workbook.active

# Load credentials
email = "dirk.ahlemeyer@gmail.com"  # Replace with your email
password = "Niederlande12345"        # Replace with your password

# Initialize WebDriver
print("Initializing WebDriver...")
driver = webdriver.Chrome()

# Maximize the browser window
print("Maximizing browser window...")
driver.maximize_window()
wait = WebDriverWait(driver, 20)

# Navigate to LeadsGorilla
print("Navigating to LeadsGorilla...")
driver.get("https://app.leadsgorilla.io/search")

time.sleep(2)

# Login process
print("Entering credentials...")
driver.find_element(By.XPATH, '//*[@id="user-name"]').send_keys(email)
driver.find_element(By.XPATH, '//*[@id="user-password"]').send_keys(password)
time.sleep(1)
print("Clicking login button...")
driver.find_element(By.XPATH, '/html/body/div/section/div/div/div/div[2]/div[2]/form/div[2]/button').click()

# Wait for the page to load after login
print("Waiting for login to complete...")
time.sleep(3)

# Iterate through each row in the Excel sheet
print("Starting to iterate through leads...")
for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, values_only=True):
    keyword, location = row  # Read keyword and location
    print(f"Processing lead: Keyword: {keyword}, Location: {location}")

    # Enter keyword and location
    print("Entering keyword...")
    driver.find_element(By.XPATH, '//*[@id="keyword-input"]').clear()
    driver.find_element(By.XPATH, '//*[@id="keyword-input"]').send_keys(keyword)

    location_input = driver.find_element(By.XPATH, '//*[@id="location"]')
    location_input.clear()
    print("Entering location...")
    location_input.send_keys(location)
    time.sleep(2)

    # Send Down Arrow and Enter keys to the location input field
    print("Sending Down arrow and Enter keys to location input...")
    location_input.send_keys(Keys.DOWN)  # Send the Down arrow key
    location_input.send_keys(Keys.ENTER)  # Send the Enter key

    time.sleep(2)
    driver.find_element(By.XPATH, '//*[@id="search"]').click() 
    time.sleep(2)
    # Click the search button
    print("Clicking the search button...")
    driver.find_element(By.XPATH, '//*[@id="search-leads"]').click()

    # Wait for results to load
    print("Waiting for search results to load...")
    time.sleep(5)
    
    if driver.find_element(By.XPATH, '/html/body/div[2]/div/div[3]/section/div/div[1]/div[6]/div[1]/p/span').text == 'LEAD RESULTS: 0':
        continue

    # Click "More" button until it's no longer present
    print("Clicking 'More' button to load more results if available...")
    while True:
        try:
            time.sleep(10)
            more_button=WebDriverWait(driver, 20).until(EC.element_to_be_clickable((By.XPATH, '//*[@id="more-btn"]')))
            time.sleep(1)
            driver.execute_script("arguments[0].scrollIntoView(true);", more_button)
            time.sleep(1)
            more_button.click()
            time.sleep(15)  # Add a slight delay to avoid overloading the server
            print("Clicked 'More' button.")
        except:
            time.sleep(3)
            print("No more results to load.")
            break

    # Select all leads
    print("Selecting all leads...")
    time.sleep(15)
    # Check if the 'Select All' button is found and its attributes
    # try:
    #     select_all_button = driver.find_element(By.XPATH, '//*[@id="select-all"]')
    #     driver.execute_script("arguments[0].scrollIntoView(true);", select_all_button)
    #     time.sleep(1)

    #     print(f"Found 'Select All' button: {select_all_button.is_displayed()}, {select_all_button.is_enabled()}")
    # except Exception as e:
    #     print(f"Error finding 'Select All' button: {str(e)}")
        
    # try:
    #     select_all_button = wait.until(EC.element_to_be_clickable((By.XPATH, '//*[@id="select-all"]')))
    #     driver.execute_script("arguments[0].scrollIntoView(true);", select_all_button)
    #     time.sleep(1)
    # except Exception as e:
    #     print(f"Error finding 'Select All' button: {str(e)}")
        
    select_all_button = wait.until(EC.element_to_be_clickable((By.XPATH, '//*[@id="select-all"]')))
    driver.execute_script("arguments[0].scrollIntoView(true);", select_all_button)
    time.sleep(1)
    select_all_button.click()
    time.sleep(5)

    # Export leads
    print("Clicking export button...")
    driver.find_element(By.XPATH, '//*[@id="export-leads-by"]').click()
    time.sleep(2)
    print("Selecting export format...")
    driver.find_element(By.XPATH, '//*[@id="export-leads-by"]/option[2]').click()
    
    

    # Wait for export to process
    print("Waiting for export to complete...")
    time.sleep(20)
    # Example: After performing some action, refresh the page
    print("Refreshing the page...")
    driver.refresh()
    time.sleep(2)  # Optionally wait for the page to load after refresh


# Close the browser
print("Closing the browser...")
driver.quit()
